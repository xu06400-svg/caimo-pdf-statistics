import os,re,sqlite3,io
from datetime import datetime
from pathlib import Path
from flask import Flask,request,jsonify,render_template,send_file
import fitz
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter

BASE=Path(__file__).resolve().parent
app=Flask(__name__)
DB=Path(os.environ.get('DB_PATH',str(BASE/'data.db')))
FIELDS=['date','document','line_no','material_code','item_name','specification','quantity','unit','production_no','sales_order','source_file']

def db():
 c=sqlite3.connect(DB);c.row_factory=sqlite3.Row
 c.execute('CREATE TABLE IF NOT EXISTS materials(id INTEGER PRIMARY KEY AUTOINCREMENT, date TEXT, document TEXT, line_no INTEGER, material_code TEXT, item_name TEXT, specification TEXT, quantity REAL, unit TEXT, production_no TEXT, sales_order TEXT, source_file TEXT)')
 c.execute('CREATE INDEX IF NOT EXISTS idx_date ON materials(date)')
 c.execute('CREATE INDEX IF NOT EXISTS idx_order ON materials(sales_order)')
 c.commit();return c

def compact(v):return re.sub(r'\s+','',v or '')
def extract(file_data, filename):
    """Read grid-based transfer requests using word coordinates, not PDF text blocks."""
    pdf = fitz.open(stream=file_data, filetype='pdf')
    results, warnings = [], []
    doc_number, doc_date = '', ''
    # Column boundaries are from the issued Chinese transfer-request PDF template.
    columns = {
        'material_code': (49, 94), 'item_name': (94, 135),
        'specification': (135, 174), 'quantity': (208, 251),
        'unit': (313, 338), 'production_no': (338, 375),
        'sales_order': (375, 416),
    }
    for page_i, page in enumerate(pdf):
        text = page.get_text()
        match_no = re.search(r'JT\d+Tra\d+', text)
        match_date = re.search(r'20\d{2}-\d{2}-\d{2}', text)
        if match_no: doc_number = match_no.group()
        if match_date: doc_date = match_date.group()
        words = page.get_text('words')
        starts = sorted([(w[1], int(w[4])) for w in words
                         if w[0] < 50 and w[1] > 205
                         and re.fullmatch(r'\d{1,3}', w[4])
                         and int(w[4]) % 10 == 0], key=lambda x: x[0])
        for i, (y, line_no) in enumerate(starts):
            # A row is 52.5pt tall; each line number sits near its middle.
            low = y - 27
            high = (starts[i + 1][0] - 27) if i + 1 < len(starts) else y + 27
            def cell(field):
                left, right = columns[field]
                found = [w for w in words if left <= (w[0] + w[2])/2 < right
                         and low <= (w[1] + w[3])/2 < high]
                found.sort(key=lambda w: (round(w[1] / 3), w[0]))
                return ''.join(w[4] for w in found).strip()
            material = cell('material_code')
            name = cell('item_name')
            raw_qty = cell('quantity')
            numeric = re.sub(r'[^0-9.,]', '', raw_qty).replace(',', '')
            try:
                if not re.fullmatch(r'\d+(?:\.\d+)?', numeric): raise ValueError()
                quantity = float(numeric)
            except ValueError:
                warnings.append(f'{filename} 第{page_i+1}页 行号{line_no}：数量无法识别，请核对')
                continue
            if not material or not name:
                warnings.append(f'{filename} 第{page_i+1}页 行号{line_no}：料号或品名缺失，请核对')
                continue
            results.append(dict(date=doc_date, document=doc_number,
                                line_no=line_no, material_code=material,
                                item_name=name, specification=cell('specification'),
                                quantity=quantity, unit=cell('unit') or '米',
                                production_no=cell('production_no'),
                                sales_order=cell('sales_order'),source_file=filename))
    return results, warnings

@app.route('/')
def home():return render_template('index.html')
@app.route('/api/upload',methods=['POST'])
def upload():
 fs=request.files.getlist('files');added=0;warnings=[];files=[]
 c=db()
 for f in fs:
  if not f.filename.lower().endswith('.pdf'):
   warnings.append(f'{f.filename}：只支持 PDF');continue
  try:
   rows,w=extract(f.read(),os.path.basename(f.filename));warnings+=w
   if not rows:warnings.append(f'{f.filename}：未识别到领料明细');continue
   # same filename re-upload replaces prior records to prevent double counting
   c.execute('DELETE FROM materials WHERE source_file=?',(os.path.basename(f.filename),))
   c.executemany('INSERT INTO materials('+','.join(FIELDS)+') VALUES ('+','.join('?' for _ in FIELDS)+')',[[r[k] for k in FIELDS] for r in rows]);c.commit()
   added+=len(rows);files.append(f.filename)
  except Exception as ex:warnings.append(f'{f.filename}：解析失败 {str(ex)}')
 c.close();return jsonify(added=added,files=files,warnings=warnings)

def query_rows():
 f=request.args;terms=[];args=[]
 for field,param,op in [('date','start','>='),('date','end','<=')]:
  if f.get(param):terms.append(f'{field}{op}?');args.append(f[param])
 for field,param in [('sales_order','order'),('material_code','material'),('document','document')]:
  if f.get(param):terms.append(f'{field} LIKE ?');args.append('%'+f[param].strip()+'%')
 q='SELECT * FROM materials'+(' WHERE '+' AND '.join(terms) if terms else '')+' ORDER BY date DESC,document,line_no,id';c=db();rows=[dict(r) for r in c.execute(q,args)];c.close();return rows
@app.route('/api/rows')
def rows():
 r=query_rows();return jsonify(rows=r,count=len(r),total=round(sum(x['quantity'] for x in r),6),documents=len({x['document'] for x in r}),orders=len({x['sales_order'] for x in r if x['sales_order']}))
@app.route('/api/export')
def export():
 rows=query_rows();wb=Workbook();ws=wb.active;ws.title='PDF领料明细';heads=['单据日期','调拨单号','行号','料号','品名','规格','申请数量','单位','生产订单号','销售订单号','来源PDF'];ws.append(heads)
 for r in rows:ws.append([r[k] for k in FIELDS]);ws.cell(ws.max_row,7).number_format='#,##0.000000'
 for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='194E74');cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(horizontal='center')
 for i,width in enumerate([15,26,9,20,34,22,18,10,24,22,42],1):ws.column_dimensions[get_column_letter(i)].width=width
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 ws2=wb.create_sheet('按订单汇总');ws2.append(['销售订单号','领料笔数','领料合计（米）'])
 totals={}
 for r in rows:
  key=r['sales_order'] or '(未填写订单号)';v=totals.setdefault(key,[0,0]);v[0]+=1;v[1]+=r['quantity']
 for k,v in sorted(totals.items()):ws2.append([k,v[0],round(v[1],6)])
 for c in ws2[1]:c.fill=PatternFill('solid',fgColor='194E74');c.font=Font(bold=True,color='FFFFFF')
 for col,w in zip('ABC',[26,18,25]):ws2.column_dimensions[col].width=w
 ws2.freeze_panes='A2';ws2.auto_filter.ref=ws2.dimensions
 b=io.BytesIO();wb.save(b);b.seek(0)
 return send_file(b,as_attachment=True,download_name='彩膜PDF领料统计.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
@app.route('/api/files')
def files():
 c=db();a=[dict(r) for r in c.execute('SELECT source_file,COUNT(*) AS count FROM materials GROUP BY source_file ORDER BY source_file')];c.close();return jsonify(files=a)
@app.route('/api/files/delete',methods=['POST'])
def delete_file():
 name=(request.json or {}).get('name');c=db();c.execute('DELETE FROM materials WHERE source_file=?',(name,));c.commit();c.close();return jsonify(ok=True)
if __name__=='__main__':
 print('浏览器打开 http://127.0.0.1:5050')
 app.run(host='127.0.0.1',port=5050,debug=False)
